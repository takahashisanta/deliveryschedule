"""
PSD Driver Schedule – Workbook Generator
Run once to create PSD_Driver_Schedule.xlsx with all required sheets.
Then import DriverSchedule.bas into the VBA editor (see README.md).
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

# ── colour palette ────────────────────────────────────────────────────────────
CLR_HEADER_DARK  = "1F3864"   # dark navy  – section headers
CLR_HEADER_MED   = "2F5496"   # mid blue   – column headers
CLR_TRUCK        = "D6E4F0"   # light blue – truck rows
CLR_VAN          = "E2EFDA"   # light green – van rows
CLR_RULES_HDR    = "F4B942"   # amber      – rules table headers
CLR_WHITE        = "FFFFFF"
CLR_LIGHT_GREY   = "F2F2F2"

# ── thin border helper ─────────────────────────────────────────────────────────
def thin_border():
    s = Side(style="thin")
    return Border(left=s, right=s, top=s, bottom=s)

def header_font(white=True):
    return Font(name="Calibri", bold=True, size=11,
                color=CLR_WHITE if white else "000000")

def cell_fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def style_header_cell(cell, text, bg=CLR_HEADER_MED, white_text=True,
                       bold=True, size=11, wrap=False):
    cell.value = text
    cell.font  = Font(name="Calibri", bold=bold, size=size,
                      color=CLR_WHITE if white_text else "000000")
    cell.fill  = cell_fill(bg)
    cell.alignment = Alignment(horizontal="center", vertical="center",
                                wrap_text=wrap)
    cell.border = thin_border()

def style_data_cell(cell, text="", bold=False, bg=CLR_WHITE, align="left"):
    if text != "":
        cell.value = text
    cell.font  = Font(name="Calibri", bold=bold, size=10)
    cell.fill  = cell_fill(bg)
    cell.alignment = Alignment(horizontal=align, vertical="center")
    cell.border = thin_border()


# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 1 – DRIVER_RULES
# ═══════════════════════════════════════════════════════════════════════════════
def build_driver_rules(ws):
    ws.sheet_view.showGridLines = False

    # ── Section title ──────────────────────────────────────────────────────────
    ws.merge_cells("A1:G1")
    c = ws["A1"]
    c.value = "DRIVER RULES – PSD Delivery Scheduling"
    c.font  = Font(name="Calibri", bold=True, size=14, color=CLR_WHITE)
    c.fill  = cell_fill(CLR_HEADER_DARK)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # ── TABLE 1: Day → Area mapping ────────────────────────────────────────────
    ws.merge_cells("A3:G3")
    c = ws["A3"]
    c.value = "TABLE 1 – DAY / AREA SCHEDULE"
    c.font  = Font(name="Calibri", bold=True, size=11, color=CLR_WHITE)
    c.fill  = cell_fill(CLR_RULES_HDR.replace("F4B942", "1F3864"))
    c.fill  = cell_fill("1F3864")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[3].height = 20

    headers_t1 = ["Day", "Area 1", "Area 2", "Area 3", "Area 4", "Area 5", "Area 6"]
    for col, h in enumerate(headers_t1, 1):
        cell = ws.cell(row=4, column=col)
        style_header_cell(cell, h, bg=CLR_HEADER_MED)

    day_areas = [
        ("Mon", "JFC/IDC", "SOUTH BAY", "CENTRAL LA", "SAN GABRIEL", "EAST LA",   ""),
        ("Tue", "LAX",     "NAX",       "JFC/IDC",    "CENTRAL LA",  "",           ""),
        ("Wed", "LAX",     "NAX",       "JFC/IDC",    "SAN GABRIEL", "NORTH OC",  "SOUTH LA"),
        ("Thu", "JFC/IDC", "SOUTH BAY", "CENTRAL LA", "EAST LA",     "",           ""),
        ("Fri", "LAX",     "NAX",       "WEST LA",    "CENTRAL LA",  "NORTH OC",  ""),
    ]
    for r, row_data in enumerate(day_areas, 5):
        bg = CLR_LIGHT_GREY if r % 2 == 0 else CLR_WHITE
        for col, val in enumerate(row_data, 1):
            c = ws.cell(row=r, column=col)
            style_data_cell(c, val, bold=(col == 1), bg=bg,
                             align="center" if col > 1 else "left")

    # ── TABLE 2: Customer → Vehicle Preference ─────────────────────────────────
    ws.merge_cells("A12:C12")
    c = ws["A12"]
    c.value = "TABLE 2 – CUSTOMER VEHICLE PREFERENCE"
    c.font  = Font(name="Calibri", bold=True, size=11, color=CLR_WHITE)
    c.fill  = cell_fill("1F3864")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[12].height = 20

    for col, h in enumerate(["Customer Name", "Vehicle", "Notes"], 1):
        style_header_cell(ws.cell(row=13, column=col), h, bg=CLR_HEADER_MED)

    customer_prefs = [
        # Truck-preferred
        ("Cal Hono",           "Truck",    ""),
        ("Kpac",               "Truck",    ""),
        ("Broad Leaf",         "Truck",    ""),
        ("AJ Foods LLC",       "Truck",    ""),
        ("SJ",                 "Truck",    ""),
        ("JK Trucking",        "Truck",    ""),
        ("Khee Trading",       "Truck",    ""),
        ("Horn Foods",         "Truck",    ""),
        ("Han Seafood",        "Truck",    ""),
        ("Maruzen",            "Truck",    ""),
        ("Honolulu Freight",   "Truck",    ""),
        ("World Food",         "Truck",    ""),
        ("HPL Apollo",         "Truck",    ""),
        ("Recycling",          "Truck",    ""),
        ("NEU",                "Truck",    ""),
        # Van-preferred
        ("Mitsuwa",            "Van",      "All Mitsuwa locations"),
        ("Turuhashi",          "Van",      ""),
        ("Hey Den",            "Van",      ""),
        ("CRC",                "Van",      ""),
        ("Sushi Zanmai",       "Van",      ""),
        # Flexible
        ("LA Cold",            "Flexible", ">=4 pallets → Truck"),
        ("JFC LA",             "Flexible", ">=4 pallets → Truck"),
        ("JFC IDC",            "Flexible", ">=4 pallets → Truck"),
        ("Wismettac",          "Flexible", ">=4 pallets → Truck"),
        ("Capital Meat",       "Flexible", ">=4 pallets → Truck"),
        ("Central Boeki",      "Flexible", ">=4 pallets → Truck"),
        ("NAX",                "Flexible", ">=4 pallets → Truck"),
        ("WCPM",               "Flexible", ">=4 pallets → Truck"),
    ]
    for r, (cust, veh, note) in enumerate(customer_prefs, 14):
        bg = CLR_TRUCK if veh == "Truck" else (CLR_VAN if veh == "Van" else CLR_LIGHT_GREY)
        style_data_cell(ws.cell(row=r, column=1), cust, bg=bg)
        style_data_cell(ws.cell(row=r, column=2), veh,  bg=bg, bold=True, align="center")
        style_data_cell(ws.cell(row=r, column=3), note, bg=bg)

    # ── TABLE 3: Time Priority ─────────────────────────────────────────────────
    t3_start = 14 + len(customer_prefs) + 2   # 2 blank rows gap

    ws.cell(row=t3_start, column=1).value = "TABLE 3 – TIME WINDOW PRIORITY"
    ws.cell(row=t3_start, column=1).font  = Font(name="Calibri", bold=True, size=11, color=CLR_WHITE)
    ws.cell(row=t3_start, column=1).fill  = cell_fill("1F3864")
    ws.cell(row=t3_start, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(f"A{t3_start}:C{t3_start}")
    ws.row_dimensions[t3_start].height = 20

    for col, h in enumerate(["Delivery Window Condition", "Priority", "Description"], 1):
        style_header_cell(ws.cell(row=t3_start+1, column=col), h, bg=CLR_HEADER_MED)

    time_priorities = [
        ("End time <= 08:00",   "1", "Must deliver by 8 AM"),
        ("End time 08:01-12:00","2", "Morning delivery"),
        ("End time 12:01-15:00","3", "Afternoon delivery"),
        ("No time window",      "9", "Flexible / no constraint"),
    ]
    for i, (cond, pri, desc) in enumerate(time_priorities):
        row = t3_start + 2 + i
        bg = CLR_LIGHT_GREY if i % 2 == 0 else CLR_WHITE
        style_data_cell(ws.cell(row=row, column=1), cond, bg=bg)
        style_data_cell(ws.cell(row=row, column=2), pri,  bg=bg, bold=True, align="center")
        style_data_cell(ws.cell(row=row, column=3), desc, bg=bg)

    # ── TABLE 4: Capacity Limits ───────────────────────────────────────────────
    t4_start = t3_start + 2 + len(time_priorities) + 2

    ws.cell(row=t4_start, column=1).value = "TABLE 4 – VEHICLE CAPACITY"
    ws.cell(row=t4_start, column=1).font  = Font(name="Calibri", bold=True, size=11, color=CLR_WHITE)
    ws.cell(row=t4_start, column=1).fill  = cell_fill("1F3864")
    ws.cell(row=t4_start, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(f"A{t4_start}:C{t4_start}")
    ws.row_dimensions[t4_start].height = 20

    for col, h in enumerate(["Vehicle", "Max Pallets", "Notes"], 1):
        style_header_cell(ws.cell(row=t4_start+1, column=col), h, bg=CLR_HEADER_MED)

    capacity = [
        ("Truck", "8", "7-8 pallets max load"),
        ("Van",   "3", "3 pallets max load"),
    ]
    for i, (veh, cap, note) in enumerate(capacity):
        row = t4_start + 2 + i
        bg = CLR_TRUCK if veh == "Truck" else CLR_VAN
        style_data_cell(ws.cell(row=row, column=1), veh,  bg=bg, bold=True)
        style_data_cell(ws.cell(row=row, column=2), cap,  bg=bg, bold=True, align="center")
        style_data_cell(ws.cell(row=row, column=3), note, bg=bg)

    # ── TABLE 5: MASTER DATA column map (for easy VBA configuration) ───────────
    t5_start = t4_start + 2 + len(capacity) + 2

    ws.cell(row=t5_start, column=1).value = "TABLE 5 – MASTER DATA COLUMN MAP  (adjust if sheet changes)"
    ws.cell(row=t5_start, column=1).font  = Font(name="Calibri", bold=True, size=11, color=CLR_WHITE)
    ws.cell(row=t5_start, column=1).fill  = cell_fill("1F3864")
    ws.cell(row=t5_start, column=1).alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(f"A{t5_start}:C{t5_start}")
    ws.row_dimensions[t5_start].height = 20

    for col, h in enumerate(["Field", "Column Letter", "Column Number"], 1):
        style_header_cell(ws.cell(row=t5_start+1, column=col), h, bg=CLR_HEADER_MED)

    col_map = [
        ("Request ID",                              "A",  "1"),
        ("Date (Ship Date)",                        "B",  "2"),
        ("Category",                                "C",  "3  – not used"),
        ("Rep",                                     "D",  "4  – not used"),
        ("Shipping Method  (Delivery / Pick-up)",   "E",  "5"),
        ("Delivery Window  (label – not used)",     "F",  "6  – not used"),
        ("Start  (window start time)",              "G",  "7"),
        ("Finish  (window end time)",               "H",  "8"),
        ("PO #",                                    "I",  "9  – not used"),
        ("Sold To Customer  (vehicle pref lookup)", "J",  "10"),
        ("Sold To Address",                         "K",  "11 – not used"),
        ("Ship To Customer  (shown on schedule)",   "L",  "12"),
        ("Ship To Address   (shown on schedule)",   "M",  "13"),
        ("Area",                                    "N",  "14"),
        ("Product / Item Description",              "O",  "15"),
        ("Quantity / Cases",                        "P",  "16"),
        ("Weight",                                  "Q",  "17 – not used"),
        ("Price / Unit",                            "R",  "18 – not used"),
        ("Storage Location",                        "S",  "19"),
        ("Version",                                 "T",  "20"),
        ("Color Sort",                              "U",  "21 – not used"),
        ("Pallets  ← ADD THIS COLUMN at col V",     "V",  "22"),
    ]
    for i, (field, col_l, col_n) in enumerate(col_map):
        row = t5_start + 2 + i
        bg = CLR_LIGHT_GREY if i % 2 == 0 else CLR_WHITE
        style_data_cell(ws.cell(row=row, column=1), field, bg=bg)
        style_data_cell(ws.cell(row=row, column=2), col_l, bg=bg, bold=True, align="center")
        style_data_cell(ws.cell(row=row, column=3), col_n, bg=bg, align="center")

    # ── Column widths ──────────────────────────────────────────────────────────
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 14
    ws.column_dimensions["D"].width = 14
    ws.column_dimensions["E"].width = 14
    ws.column_dimensions["F"].width = 14
    ws.column_dimensions["G"].width = 14


# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 2 – DRIVER_ENGINE  (hidden calculation table)
# ═══════════════════════════════════════════════════════════════════════════════
def build_driver_engine(ws):
    ws.sheet_view.showGridLines = True

    headers = [
        "Request ID",          # A
        "Ship Date",           # B
        "Sold To (Customer)",  # C
        "Ship To",             # D
        "Area",                # E
        "Product",             # F
        "Qty / Pallets",       # G
        "Weight",              # H
        "Storage Loc",         # I
        "Window Start",        # J
        "Window End",          # K
        "Delivery Note",       # L
        "Vehicle Pref (Raw)",  # M
        "Assigned Vehicle",    # N
        "Time Priority",       # O
        "Area Rank",           # P
        "Sort Key",            # Q
    ]
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col)
        style_header_cell(c, h, bg=CLR_HEADER_DARK)
        ws.column_dimensions[get_column_letter(col)].width = 18

    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    # Add a note row so it's clear this sheet is managed by macro
    note_cell = ws["A2"]
    note_cell.value = "⬅  This sheet is populated automatically by the GenerateSchedule macro. Do not edit manually."
    note_cell.font  = Font(name="Calibri", italic=True, size=10, color="7F7F7F")
    ws.merge_cells("A2:Q2")


# ═══════════════════════════════════════════════════════════════════════════════
# SHEET 3 – DRIVER SCHEDULE  (printable output)
# ═══════════════════════════════════════════════════════════════════════════════
def build_driver_schedule(ws):
    ws.sheet_view.showGridLines = False

    # ── Title block ────────────────────────────────────────────────────────────
    ws.merge_cells("A1:H1")
    c = ws["A1"]
    c.value = "PSD DRIVER SCHEDULE"
    c.font  = Font(name="Calibri", bold=True, size=18, color=CLR_WHITE)
    c.fill  = cell_fill(CLR_HEADER_DARK)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # ── Date selector ──────────────────────────────────────────────────────────
    ws["A2"].value = "Delivery Date:"
    ws["A2"].font  = Font(name="Calibri", bold=True, size=12)
    ws["A2"].alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[2].height = 24

    ws.merge_cells("B2:C2")
    c = ws["B2"]
    c.value = None   # user enters date here
    c.font  = Font(name="Calibri", bold=True, size=12, color="1F3864")
    c.fill  = cell_fill("FFF2CC")   # yellow – input cell
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = thin_border()
    c.number_format = "MM/DD/YYYY"

    ws["D2"].value = "← Enter date, then click Generate Schedule"
    ws["D2"].font  = Font(name="Calibri", italic=True, size=10, color="7F7F7F")
    ws["D2"].alignment = Alignment(vertical="center")

    # ── Button placeholder note (actual button added in Excel) ─────────────────
    ws.merge_cells("F2:H2")
    c = ws["F2"]
    c.value = "[ GENERATE SCHEDULE button goes here ]"
    c.font  = Font(name="Calibri", bold=True, size=10, color="7F7F7F")
    c.fill  = cell_fill("D9D9D9")
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = thin_border()

    ws.row_dimensions[3].height = 8   # spacer

    # ── TRUCK section header ───────────────────────────────────────────────────
    ws.merge_cells("A4:H4")
    c = ws["A4"]
    c.value = "🚚  TRUCK ROUTE"
    c.font  = Font(name="Calibri", bold=True, size=13, color=CLR_WHITE)
    c.fill  = cell_fill("1F4E79")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[4].height = 24

    truck_cols = ["Stop #", "Customer", "Ship To / Area", "Product",
                  "Qty/Pallets", "Window", "Storage", "Notes"]
    for col, h in enumerate(truck_cols, 1):
        style_header_cell(ws.cell(row=5, column=col), h, bg=CLR_HEADER_MED)
    ws.row_dimensions[5].height = 20

    # 8 blank truck rows (will be overwritten by macro)
    for r in range(6, 14):
        bg = CLR_TRUCK if r % 2 == 0 else CLR_WHITE
        for col in range(1, 9):
            style_data_cell(ws.cell(row=r, column=col), bg=bg)
        ws.row_dimensions[r].height = 18

    # ── spacer ─────────────────────────────────────────────────────────────────
    ws.row_dimensions[14].height = 10

    # ── VAN section header ─────────────────────────────────────────────────────
    ws.merge_cells("A15:H15")
    c = ws["A15"]
    c.value = "🚐  VAN ROUTE"
    c.font  = Font(name="Calibri", bold=True, size=13, color=CLR_WHITE)
    c.fill  = cell_fill("375623")
    c.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[15].height = 24

    van_cols = ["Stop #", "Customer", "Ship To / Area", "Product",
                "Qty/Pallets", "Window", "Storage", "Notes"]
    for col, h in enumerate(van_cols, 1):
        style_header_cell(ws.cell(row=16, column=col), h, bg="538135")
    ws.row_dimensions[16].height = 20

    for r in range(17, 22):
        bg = CLR_VAN if r % 2 == 0 else CLR_WHITE
        for col in range(1, 9):
            style_data_cell(ws.cell(row=r, column=col), bg=bg)
        ws.row_dimensions[r].height = 18

    # ── Column widths ──────────────────────────────────────────────────────────
    widths = [7, 22, 22, 22, 10, 16, 14, 20]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ── Print settings ─────────────────────────────────────────────────────────
    ws.page_setup.orientation    = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize      = ws.PAPERSIZE_LETTER
    ws.page_setup.fitToPage      = True
    ws.page_setup.fitToWidth     = 1
    ws.page_setup.fitToHeight    = 0
    ws.print_area                = "A1:H40"


# ═══════════════════════════════════════════════════════════════════════════════
# MAIN
# ═══════════════════════════════════════════════════════════════════════════════
def main():
    wb = openpyxl.Workbook()
    default = wb.active
    wb.remove(default)

    ws_rules    = wb.create_sheet("DRIVER_RULES")
    ws_engine   = wb.create_sheet("DRIVER_ENGINE")
    ws_schedule = wb.create_sheet("DRIVER_SCHEDULE")

    print("Building DRIVER_RULES …")
    build_driver_rules(ws_rules)

    print("Building DRIVER_ENGINE …")
    build_driver_engine(ws_engine)
    ws_engine.sheet_state = "hidden"

    print("Building DRIVER_SCHEDULE …")
    build_driver_schedule(ws_schedule)

    # Make DRIVER_SCHEDULE the active tab
    ws_schedule.sheet_view.tabSelected = True
    wb.active = ws_schedule

    out = "PSD_Driver_Schedule.xlsx"
    wb.save(out)
    print(f"\n✅  Saved: {out}")
    print("Next step: open the file in Excel, then import DriverSchedule.bas")
    print("See README.md for full setup instructions.")


if __name__ == "__main__":
    main()
